"""
Compare two LCA quantity sheets (mængdeopgørelser) and highlight the differences.

The two sheets must have the same column layout:
    column A = TypeCode, column B = part name (Bygningsdelsnavn),
    column C = extra description, the remaining columns = quantities/units/notes.
Headers are on row 10 and data starts on row 11 (both can be changed with options).

The result is a copy of the NEW workbook where, in the compared sheet:
    * new rows are filled green,
    * changed cells are filled yellow and get a comment with the old value,
    * a status column is added to the right of the data,
and a "Comparison" sheet listing every new, changed and removed line.

If no files are given, the single Excel file in the folder "old_LCA_sheet" is
used as the old file and the single Excel file in "new_LCA_sheet" as the new
file (both folders next to this script).

Usage:
    py compare_lca.py --old-sheet ARK --new-sheet Ark1
    py compare_lca.py --list-sheets
    py compare_lca.py OLD.xlsx NEW.xlsx --old-sheet ARK --new-sheet Ark1
    py compare_lca.py OLD.xlsx NEW.xlsx --old-sheet ARK --new-sheet Ark1 -o result.xlsx
"""

import argparse
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CODE_COL, NAME_COL = 0, 1  # 0-based positions in a row

PROJECT_DIR = Path(__file__).resolve().parent
OLD_DIR = PROJECT_DIR / "old_LCA_sheet"
NEW_DIR = PROJECT_DIR / "new_LCA_sheet"
OUTPUT_SUFFIX = " - comparison.xlsx"

FILL_NEW = PatternFill("solid", start_color="C6EFCE")
FILL_CHANGED = PatternFill("solid", start_color="FFEB9C")
FILL_REMOVED = PatternFill("solid", start_color="FFC7CE")
FILL_HEADER = PatternFill("solid", start_color="D9D9D9")


@dataclass
class Row:
    excel_row: int
    values: list

    @property
    def code(self):
        return key_text(self.values[CODE_COL])

    @property
    def name(self):
        return key_text(self.values[NAME_COL])


# ---------------------------------------------------------------- reading ---

def key_text(v):
    """Text used for matching rows: 216001 and '216001' are the same code."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def find_excel_file(folder):
    """The one Excel file in `folder` (Excel lock files and earlier reports are ignored)."""
    if not folder.is_dir():
        sys.exit(f"Folder not found: {folder}")
    files = [f for f in folder.iterdir()
             if f.suffix.lower() in (".xlsx", ".xlsm")
             and not f.name.startswith("~$") and not f.name.endswith(OUTPUT_SUFFIX)]
    if len(files) != 1:
        found = ", ".join(f.name for f in files) or "none"
        sys.exit(f"Expected exactly one Excel file in {folder}, found {len(files)}: {found}")
    return files[0]


def pick_sheet(wb, sheet_name, path):
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            sys.exit(f"Sheet '{sheet_name}' not found in {path}. Available sheets: {wb.sheetnames}")
        return wb[sheet_name]
    if len(wb.sheetnames) == 1:
        return wb.worksheets[0]
    sys.exit(f"{path} has several sheets {wb.sheetnames} - choose one with --old-sheet/--new-sheet.")


def read_sheet(path, sheet_name, header_row, start_row):
    wb = openpyxl.load_workbook(path, data_only=True)  # data_only -> values, not formulas
    ws = pick_sheet(wb, sheet_name, path)
    header_cells = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    ncols = max((i + 1 for i, h in enumerate(header_cells) if h not in (None, "")), default=0)
    if ncols < 3:
        sys.exit(f"Could not find headers on row {header_row} of '{ws.title}' in {path}.")
    headers = [" ".join(str(h).split()) if h is not None else get_column_letter(i + 1)
               for i, h in enumerate(header_cells[:ncols])]

    rows = []
    for r in range(start_row, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ncols + 1)]
        if any(normalize(v) is not None for v in values):
            rows.append(Row(r, values))
    return ws.title, headers, rows


# ------------------------------------------------------------- comparing ---

def normalize(v):
    """Make values comparable: blank strings == None, trimmed text, numbers as float."""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = "\n".join(line.strip() for line in v.strip().splitlines())
        return s or None
    return v  # dates etc.


def equal(a, b, tol):
    a, b = normalize(a), normalize(b)
    # a number stored as text in one file and as a number in the other
    if isinstance(a, float) and isinstance(b, str) or isinstance(a, str) and isinstance(b, float):
        try:
            a, b = float(a), float(b)
        except ValueError:
            return False
    if isinstance(a, float) and isinstance(b, float):
        return abs(a - b) <= tol * max(1.0, abs(a), abs(b))
    return a == b


def match_rows(old_rows, new_rows):
    """Return (pairs, new_only, removed).

    1. Rows are paired on TypeCode + name (n-th occurrence with n-th occurrence,
       since placeholder codes like '253xxx' are reused for many parts).
    2. Rows left over are paired on TypeCode alone when that code occurs exactly
       once in each sheet (= the part was renamed).
    """
    pairs = []
    old_by_key = defaultdict(list)
    for row in old_rows:
        old_by_key[(row.code, row.name)].append(row)

    unmatched_new = []
    for row in new_rows:
        bucket = old_by_key.get((row.code, row.name))
        if bucket:
            pairs.append((bucket.pop(0), row))
        else:
            unmatched_new.append(row)
    unmatched_old = [r for bucket in old_by_key.values() for r in bucket]

    old_counts = Counter(r.code for r in old_rows)
    new_counts = Counter(r.code for r in new_rows)
    old_by_code = {r.code: r for r in unmatched_old if old_counts[r.code] == 1}
    new_only = []
    for row in unmatched_new:
        if new_counts[row.code] == 1 and row.code in old_by_code:
            pairs.append((old_by_code.pop(row.code), row))
        else:
            new_only.append(row)
    matched_old = {id(o) for o, _ in pairs}
    removed = [r for r in old_rows if id(r) not in matched_old]
    return pairs, new_only, removed


def compare(old_rows, new_rows, tol):
    pairs, new_only, removed = match_rows(old_rows, new_rows)
    changes = []  # (old_row, new_row, [col indexes that differ])
    for old, new in pairs:
        ncols = max(len(old.values), len(new.values))
        get = lambda vals, i: vals[i] if i < len(vals) else None
        diff = [c for c in range(ncols) if not equal(get(old.values, c), get(new.values, c), tol)]
        if diff:
            changes.append((old, new, diff))
    changes.sort(key=lambda ch: ch[1].excel_row)
    return changes, new_only, removed


# ---------------------------------------------------------------- output ---

def fmt(v):
    return "(empty)" if normalize(v) is None else str(v)


def write_report(new_path, new_sheet, headers, changes, new_only, removed, out_path, info):
    wb = openpyxl.load_workbook(new_path)  # keep formulas and formatting of the new file
    ws = wb[new_sheet]
    ncols = len(headers)
    status_col = ws.max_column + 1
    header_row = info["header_row"]

    ws.cell(header_row, status_col, "Comparison status").font = Font(bold=True)
    ws.cell(header_row, status_col).fill = FILL_HEADER

    for row in new_only:
        for c in range(1, ncols + 1):
            ws.cell(row.excel_row, c).fill = FILL_NEW
        ws.cell(row.excel_row, status_col, "NEW").fill = FILL_NEW

    for old, new, diff in changes:
        for c in diff:
            cell = ws.cell(new.excel_row, c + 1)
            cell.fill = FILL_CHANGED
            comment = Comment(f"Old value (row {old.excel_row}):\n{fmt(old.values[c] if c < len(old.values) else None)}",
                              "LCA compare")
            comment.width, comment.height = 300, 150
            cell.comment = comment
        names = ", ".join(headers[c] for c in diff)
        ws.cell(new.excel_row, status_col, f"CHANGED: {names}").fill = FILL_CHANGED

    # ---- summary sheet
    if "Comparison" in wb.sheetnames:
        del wb["Comparison"]
    rs = wb.create_sheet("Comparison", 0)
    lines = [
        ("Old file", f"{info['old_path']}  [sheet: {info['old_sheet']}]"),
        ("New file", f"{info['new_path']}  [sheet: {info['new_sheet']}]"),
        ("New lines", len(new_only)),
        ("Changed lines", len(changes)),
        ("Removed lines", len(removed)),
    ]
    for i, (k, v) in enumerate(lines, start=1):
        rs.cell(i, 1, k).font = Font(bold=True)
        rs.cell(i, 2, v)

    top = len(lines) + 2
    cols = ["Status", "Old row", "New row", headers[CODE_COL], headers[NAME_COL],
            "Changed column", "Old value", "New value"]
    for c, h in enumerate(cols, start=1):
        cell = rs.cell(top, c, h)
        cell.font = Font(bold=True)
        cell.fill = FILL_HEADER

    def add(status, fill, old, new, column="", old_val="", new_val=""):
        nonlocal r
        src = new or old
        values = [status, old.excel_row if old else "", new.excel_row if new else "",
                  src.values[CODE_COL], src.values[NAME_COL], column, old_val, new_val]
        for c, v in enumerate(values, start=1):
            cell = rs.cell(r, c, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        rs.cell(r, 1).fill = fill
        if new:
            rs.cell(r, 3).hyperlink = f"#'{new_sheet}'!A{new.excel_row}"
            rs.cell(r, 3).font = Font(color="0563C1", underline="single")
        r += 1

    r = top + 1
    events = [(n.excel_row, 0, "new", n) for n in new_only] + \
             [(ch[1].excel_row, 1, "changed", ch) for ch in changes]
    for _, _, kind, item in sorted(events, key=lambda e: e[:2]):
        if kind == "new":
            add("NEW", FILL_NEW, None, item, "(entire line)")
        else:
            old, new, diff = item
            for c in diff:
                add("CHANGED", FILL_CHANGED, old, new, headers[c],
                    fmt(old.values[c] if c < len(old.values) else None),
                    fmt(new.values[c] if c < len(new.values) else None))
    for old in removed:
        add("REMOVED", FILL_REMOVED, old, None, "(entire line)")

    for c, width in enumerate([12, 9, 9, 14, 45, 25, 45, 45], start=1):
        rs.column_dimensions[get_column_letter(c)].width = width
    rs.freeze_panes = rs.cell(top + 1, 1)
    rs.auto_filter.ref = f"A{top}:{get_column_letter(len(cols))}{max(r - 1, top)}"

    wb.active = 0
    wb.save(out_path)


# ------------------------------------------------------------------ main ---

def main():
    p = argparse.ArgumentParser(description="Compare two LCA Excel sheets and highlight differences.")
    p.add_argument("old", nargs="?", help=f"old Excel file (default: the one file in {OLD_DIR.name})")
    p.add_argument("new", nargs="?", help=f"new Excel file (default: the one file in {NEW_DIR.name})")
    p.add_argument("--old-sheet", help="sheet to read in the old file (needed if it has several sheets)")
    p.add_argument("--new-sheet", help="sheet to read in the new file (needed if it has several sheets)")
    p.add_argument("-o", "--output",
                   help=f"result file (default: '<new file name>{OUTPUT_SUFFIX}' in the project folder)")
    p.add_argument("--header-row", type=int, default=10, help="row with column headers (default 10)")
    p.add_argument("--start-row", type=int, default=11, help="first data row (default 11)")
    p.add_argument("--tolerance", type=float, default=1e-9,
                   help="relative tolerance when comparing numbers (default 1e-9)")
    p.add_argument("--list-sheets", action="store_true", help="only list the sheets in the given file(s)")
    args = p.parse_args()

    if args.old and not args.new:
        p.error("give both an old and a new file, or neither to use the default folders")
    args.old = args.old or find_excel_file(OLD_DIR)
    args.new = args.new or find_excel_file(NEW_DIR)

    if args.list_sheets:
        for f in (args.old, args.new):
            print(f"{f}: {openpyxl.load_workbook(f, read_only=True).sheetnames}")
        return

    old_sheet, old_headers, old_rows = read_sheet(args.old, args.old_sheet, args.header_row, args.start_row)
    new_sheet, new_headers, new_rows = read_sheet(args.new, args.new_sheet, args.header_row, args.start_row)

    if old_headers != new_headers:
        print("WARNING: the column headers differ - check that the right sheets were chosen.")
        for i in range(max(len(old_headers), len(new_headers))):
            o = old_headers[i] if i < len(old_headers) else "-"
            n = new_headers[i] if i < len(new_headers) else "-"
            if o != n:
                print(f"  column {get_column_letter(i + 1)}: old '{o}'  new '{n}'")

    changes, new_only, removed = compare(old_rows, new_rows, args.tolerance)

    out = Path(args.output) if args.output else PROJECT_DIR / (Path(args.new).stem + OUTPUT_SUFFIX)
    info = dict(old_path=args.old, new_path=args.new, old_sheet=old_sheet, new_sheet=new_sheet,
                header_row=args.header_row)
    write_report(args.new, new_sheet, new_headers, changes, new_only, removed, out, info)

    print(f"Old: {args.old} [{old_sheet}]  {len(old_rows)} lines")
    print(f"New: {args.new} [{new_sheet}]  {len(new_rows)} lines")
    print(f"  new lines:     {len(new_only)}")
    print(f"  changed lines: {len(changes)}")
    print(f"  removed lines: {len(removed)}")
    print(f"Result written to: {out}")


if __name__ == "__main__":
    main()
